"""Unit tests for all file-type handlers."""

import io

import docx
import fitz  # PyMuPDF
import pytest
from core.handlers import (
    CsvHandler,
    DocxHandler,
    PDFHandler,
    TextHandler,
    supported_formats,
)

# ── TextHandler ──────────────────────────────────────────────


@pytest.mark.anyio
async def test_text_handler_plain():
    handler = TextHandler()
    result = await handler.parse(b"Hello World", "test.txt")
    assert result.text == "Hello World"
    assert len(result.pages) == 1
    assert result.pages[0]["page"] == 1
    assert result.pages[0]["text"] == "Hello World"
    assert result.images == []


@pytest.mark.anyio
async def test_text_handler_markdown():
    handler = TextHandler()
    data = b"# Title\n\nSome paragraph"
    result = await handler.parse(data, "readme.md")
    assert "# Title" in result.text
    assert result.pages[0]["page"] == 1


@pytest.mark.anyio
async def test_text_handler_idempotent():
    handler = TextHandler()
    data = b"same input"
    r1 = await handler.parse(data, "a.txt")
    r2 = await handler.parse(data, "a.txt")
    assert r1.text == r2.text
    assert r1.pages == r2.pages


# ── CsvHandler ───────────────────────────────────────────────


@pytest.mark.anyio
async def test_csv_handler_basic():
    handler = CsvHandler()
    data = b"Name,Age\nAlice,30\nBob,25"
    result = await handler.parse(data, "test.csv")
    assert "Name | Age" in result.text
    assert "---" in result.text
    assert "Alice | 30" in result.text
    assert result.metadata["row_count"] == 2
    assert result.metadata["columns"] == ["Name", "Age"]


@pytest.mark.anyio
async def test_csv_handler_pages():
    handler = CsvHandler()
    data = b"H1,H2\nA,B\nC,D"
    result = await handler.parse(data, "data.csv")
    assert result.pages[0]["page"] == 1
    assert result.pages[1]["page"] == 2


@pytest.mark.anyio
async def test_csv_handler_empty():
    handler = CsvHandler()
    result = await handler.parse(b"", "empty.csv")
    assert result.text == ""
    assert result.pages == []


@pytest.mark.anyio
async def test_csv_handler_bom():
    handler = CsvHandler()
    data = b"\xef\xbb\xbfName,Value\nX,1"
    result = await handler.parse(data, "bom.csv")
    assert "Name" in result.text


@pytest.mark.anyio
async def test_csv_handler_idempotent():
    handler = CsvHandler()
    data = b"A,B\n1,2"
    r1 = await handler.parse(data, "t.csv")
    r2 = await handler.parse(data, "t.csv")
    assert r1.text == r2.text


# ── DocxHandler ──────────────────────────────────────────────


def _make_docx_bytes(paragraphs: list[str], table_data: list[list[str]] | None = None) -> bytes:
    """Create a minimal .docx file in memory."""
    doc = docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for i, row in enumerate(table_data):
            for j, cell_text in enumerate(row):
                table.rows[i].cells[j].text = cell_text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.anyio
async def test_docx_handler_paragraphs():
    handler = DocxHandler()
    data = _make_docx_bytes(["First paragraph", "Second paragraph"])
    result = await handler.parse(data, "test.docx")
    assert "First paragraph" in result.text
    assert "Second paragraph" in result.text
    assert result.pages[0]["page"] == 1


@pytest.mark.anyio
async def test_docx_handler_table():
    handler = DocxHandler()
    data = _make_docx_bytes(
        ["Before table"],
        table_data=[["Col1", "Col2"], ["A", "B"]],
    )
    result = await handler.parse(data, "table.docx")
    assert "| Col1 | Col2 |" in result.text
    assert "| --- | --- |" in result.text
    assert "| A | B |" in result.text


@pytest.mark.anyio
async def test_docx_handler_empty_paragraphs_skipped():
    handler = DocxHandler()
    data = _make_docx_bytes(["Content", "", "  ", "More content"])
    result = await handler.parse(data, "gaps.docx")
    assert "Content" in result.text
    assert "More content" in result.text


@pytest.mark.anyio
async def test_docx_handler_idempotent():
    handler = DocxHandler()
    data = _make_docx_bytes(["Stable"])
    r1 = await handler.parse(data, "s.docx")
    r2 = await handler.parse(data, "s.docx")
    assert r1.text == r2.text


# ── PDFHandler (PyMuPDF fallback) ────────────────────────────


def _make_pdf_bytes(page_texts: list[str]) -> bytes:
    """Create a minimal PDF with one text block per page."""
    doc = fitz.open()
    for text in page_texts:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.mark.anyio
async def test_pdf_handler_pymupdf_single_page():
    handler = PDFHandler()
    data = _make_pdf_bytes(["Page one content"])
    result = await handler.parse(data, "test.pdf")
    assert "Page one content" in result.text
    assert len(result.pages) == 1
    assert result.pages[0]["page"] == 1
    assert result.metadata["parser"] == "pymupdf"
    assert result.metadata["total_pages"] == 1


@pytest.mark.anyio
async def test_pdf_handler_pymupdf_multi_page():
    handler = PDFHandler()
    data = _make_pdf_bytes(["First", "Second", "Third"])
    result = await handler.parse(data, "multi.pdf")
    assert len(result.pages) == 3
    assert result.pages[0]["page"] == 1
    assert result.pages[1]["page"] == 2
    assert result.pages[2]["page"] == 3
    assert result.metadata["total_pages"] == 3


@pytest.mark.anyio
async def test_pdf_handler_pymupdf_idempotent():
    handler = PDFHandler()
    data = _make_pdf_bytes(["Stable content"])
    r1 = await handler.parse(data, "s.pdf")
    r2 = await handler.parse(data, "s.pdf")
    assert r1.text == r2.text
    assert r1.pages == r2.pages


# ── Registry & supported_formats ─────────────────────────────


def test_supported_formats():
    fmts = supported_formats()
    assert fmts == [
        ".cnc", ".csv", ".docx", ".md", ".mpf", ".nc", ".pdf", ".spf", ".txt"
    ]


# ── CNC / G-Code handler ─────────────────────────────────────

_CNC_BYTES = (
    b"%_N_112217126_MPF\r\n;11221.7126\r\n"
    b"T1 M16 ;DM=9 VHMI-BOHRER\r\n"
    b"G0 G54 X1155 Y28 S10900 F3000 M3 M7 M8\r\nMCALL CYCLE82(5,0,5,-9.5,,0.1)\r\nM30\r\n"
)


@pytest.mark.anyio
async def test_cnc_extensionless_sniffed_and_parsed():
    """A file WITHOUT extension is recognised as CNC by content."""
    from core.cleaner import clean

    doc = await clean(_CNC_BYTES, "112217126")  # no extension
    assert doc.metadata["format"] == "cnc"
    assert "VHMI-BOHRER" in doc.text


@pytest.mark.anyio
async def test_cnc_by_extension():
    from core.cleaner import clean

    doc = await clean(_CNC_BYTES, "prog.mpf")
    assert doc.metadata["format"] == "cnc"


def test_cnc_decode_cp1252_umlaut():
    from core.handlers import decode_cnc_bytes

    # 0xE4 = ä in cp1252 (invalid UTF-8) must not crash and must round-trip
    assert decode_cnc_bytes(b"Fr\xe4sen") == "Fräsen"


# ── Unsupported format via cleaner ───────────────────────────


@pytest.mark.anyio
async def test_unsupported_format_raises():
    from core.cleaner import clean

    # Non-CNC content with an unknown extension still raises.
    with pytest.raises(ValueError, match="Unsupported file format"):
        await clean(b"data", "file.xyz")


@pytest.mark.anyio
async def test_unsupported_format_lists_supported():
    from core.cleaner import clean

    with pytest.raises(ValueError, match=r"\.pdf"):
        await clean(b"data", "file.xyz")
